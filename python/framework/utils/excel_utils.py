from pathlib import Path

from openpyxl import Workbook, load_workbook


def read_sheet(file_path, sheet_name):
    path = Path(file_path)
    if not path.exists():
        raise RuntimeError(f"Excel file not found: {path}")

    wb = load_workbook(path, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise RuntimeError(f"Sheet not found: {sheet_name}")

        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            return []

        header = rows[0]
        cols = len(header)
        data = []

        for row in rows[1:]:
            row_values = []
            for idx in range(cols):
                cell = row[idx] if idx < len(row) else None
                row_values.append("" if cell is None else str(cell))
            data.append(row_values)

        return data
    finally:
        wb.close()


def ensure_sample_login_workbook(file_path):
    path = Path(file_path)
    if path.exists():
        return

    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Login"

    header = ["testName", "username", "password", "expectedContains", "run"]
    sheet.append(header)

    sheet.append(
        ["valid_login", "tomsmith", "SuperSecretPassword!", "You logged into a secure area!", "Y"]
    )
    sheet.append(["invalid_login", "wrong", "wrong", "Your username is invalid!", "Y"])

    wb.save(path)
